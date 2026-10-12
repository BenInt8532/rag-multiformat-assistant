"""
Загрузчики документов разных форматов.
Каждый загрузчик читает файл и приводит его к общему виду: текст и таблицы.

Общий вид текста:
- абзацы разделены пустой строкой;
- заголовок - отдельный абзац, который начинается с "# ";
- пункты списков - строки, которые начинаются с "- " или с номера ("1. ").
"""

import csv
import io
import re
import statistics
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import List

# Расширения файлов по видам
TEXT_EXTENSIONS = {".txt", ".md"}
HTML_EXTENSIONS = {".html", ".htm"}
CODE_EXTENSIONS = {".py"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | HTML_EXTENSIONS | CODE_EXTENSIONS | {".pdf", ".docx", ".csv", ".xlsx"}

# Части страницы, которые не относятся к содержанию
HTML_JUNK_TAGS = ["script", "style", "noscript", "nav", "aside", "button", "svg"]
# Шапка и подвал убираются, только если это шапка и подвал всей страницы, а не статьи
HTML_PAGE_FRAME_TAGS = ["header", "footer"]
HTML_CONTENT_PARENTS = ["article", "main", "section"]
HTML_BLOCK_TAGS = [
    "h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "dt", "dd", "table", "pre", "blockquote",
    "div", "section", "article", "main", "body", "figcaption", "summary", "label"
]

# Пункт списка: номер ("1.", "2)") или маркер в начале строки. Общее правило для загрузчиков и нарезки
LIST_ITEM = re.compile(r"^\s*(\d+[.)]|[-•*–—])\s+")

# Номер страницы с подписью: "Стр. 5", "Страница 5 из 12", "Page 5 of 12"
PAGE_LABEL = re.compile(r"^\s*(стр\.?|страница|page)\s*\d+(\s*(из|of|/)\s*\d+)?\s*$", re.IGNORECASE)
# Номер страницы без подписи: "5", "- 5 -"
PAGE_BARE_NUMBER = re.compile(r"^\s*[-–—]?\s*(\d+)\s*[-–—]?\s*$")

# Сколько строк у верхнего и нижнего края страницы PDF проверяется на колонтитул
PDF_EDGE_LINES = 2
# Пункт списка в PDF: номер, дефис или точка-маркер. Длинное тире в начале строки -
# обычно продолжение предложения ("ЛидПоток" / "— это сеть"), а не пункт
PDF_LIST_ITEM = re.compile(r"^\s*(\d+[.)]|[-•])\s+")

# Частицы, которые пишутся через дефис: "какой-" + "нибудь" - не перенос слова
HYPHEN_PARTICLES = ("то", "либо", "нибудь", "таки", "ка")

# Кодировки текстовых файлов: utf-8 (с меткой и без) и кодировка русского Excel
TEXT_ENCODINGS = ("utf-8-sig", "cp1251")


@dataclass
class Table:
    """Таблица документа: название, заголовки столбцов и строки."""
    name: str
    header: List[str]
    rows: List[List[str]]


@dataclass
class Document:
    """Документ в общем виде: текст и таблицы, найденные в файле."""
    name: str
    file_type: str
    text: str = ""
    tables: List[Table] = field(default_factory=list)


def list_documents(data_dir: Path) -> List[Path]:
    """
    Список файлов базы знаний.

    Args:
        data_dir: папка с документами

    Returns:
        пути к файлам поддерживаемых форматов, по алфавиту
    """
    if not data_dir.exists():
        raise FileNotFoundError(f"Папка с документами {data_dir} не найдена")

    return sorted(
        path for path in data_dir.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )


def load_document(path: Path, clean: bool = True) -> Document:
    """
    Загрузка одного документа.

    Args:
        path: путь к файлу
        clean: очищать ли текст (служебные части страниц, переносы, лишние пробелы)

    Returns:
        документ в общем виде
    """
    extension = path.suffix.lower()
    document = Document(name=path.name, file_type=extension.lstrip("."))

    if extension in CODE_EXTENSIONS:
        # Код не очищается: отступы и пустые строки в нём значимы
        document.text = _read_text(path)
        return document

    if extension in TEXT_EXTENSIONS:
        document.text = _read_text(path)
    elif extension in HTML_EXTENSIONS:
        _load_html(path, document, clean)
    elif extension == ".pdf":
        _load_pdf(path, document, clean)
    elif extension == ".docx":
        _load_docx(path, document)
    elif extension == ".csv":
        _load_csv(path, document)
    elif extension == ".xlsx":
        _load_xlsx(path, document)
    else:
        raise ValueError(f"Формат {extension} не поддерживается: {path.name}")

    if clean:
        document.text = normalize_whitespace(document.text)
        for table in document.tables:
            table.header = [normalize_whitespace(cell) for cell in table.header]
            table.rows = [[normalize_whitespace(cell) for cell in row] for row in table.rows]

    return document


def normalize_whitespace(text: str) -> str:
    """
    Нормализация пробелов и переносов строк.

    Args:
        text: исходный текст

    Returns:
        текст, в котором пробелы идут по одному, а пустые строки - не больше одной подряд
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Табуляция и неразрывные пробелы для модели - лишние символы
    text = re.sub(r"[ \t ​]+", " ", text)
    text = re.sub(r" ?\n ?", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _read_text(path: Path) -> str:
    """
    Чтение текстового файла. Сначала как utf-8 (utf-8-sig убирает служебную метку
    в начале файла), при ошибке - как cp1251: так сохраняет CSV русский Excel.
    """
    data = path.read_bytes()

    # Метка в начале файла однозначно указывает на UTF-16 ("Текст Юникод" из Excel)
    encodings = ("utf-16",) if data[:2] in (b"\xff\xfe", b"\xfe\xff") else TEXT_ENCODINGS

    for encoding in encodings:
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        # Нулевые символы в тексте - признак того, что кодировка угадана неверно
        if "\x00" in text:
            continue
        # Перевод строки Windows приводится к общему виду всегда: это не очистка текста,
        # без этого абзацы не разделяются по пустой строке
        return text.replace("\r\n", "\n").replace("\r", "\n")

    raise ValueError(f"не удалось определить кодировку (пробовали {', '.join(encodings)})")


def _require(module_name: str, package: str, extension: str):
    """Подключение библиотеки формата с понятной ошибкой, если она не установлена."""
    try:
        return __import__(module_name)
    except ImportError:
        raise ImportError(
            f"Для файлов {extension} нужна библиотека {package}: pip install {package}"
        )


def _cell_text(value) -> str:
    """Значение ячейки таблицы как строка; пустая ячейка - пустая строка."""
    return "" if value is None else str(value).strip()


def _add_table(document: Document, name: str, rows: List[List[str]]):
    """
    Добавление таблицы в документ. Первая непустая строка считается заголовками столбцов.

    Args:
        document: документ, в который добавляется таблица
        name: название таблицы
        rows: все строки таблицы, включая заголовки
    """
    rows = [row for row in rows if any(cell for cell in row)]
    if len(rows) < 2:
        return

    # Все строки выравниваются по самой длинной: лишняя ячейка не должна потеряться
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]

    header, body = rows[0], rows[1:]
    # Столбец без названия получает номер, чтобы значение не осталось без подписи
    header = [cell or f"столбец {i}" for i, cell in enumerate(header, 1)]
    document.tables.append(Table(name=name, header=header, rows=body))


def _load_html(path: Path, document: Document, clean: bool):
    """Загрузка HTML-страницы: заголовки, абзацы, списки и таблицы без разметки."""
    bs4 = _require("bs4", "beautifulsoup4", ".html")
    # Странице передаются байты: кодировку BeautifulSoup берёт из её собственной метки charset
    soup = bs4.BeautifulSoup(path.read_bytes(), "html.parser")

    if not clean:
        # Без очистки: весь текст страницы подряд, вместе с меню, шапкой и подвалом
        document.text = soup.get_text("\n").replace("\r\n", "\n").replace("\r", "\n")
        return

    for tag in soup.find_all(HTML_JUNK_TAGS):
        tag.decompose()
    for tag in soup.find_all(HTML_PAGE_FRAME_TAGS):
        if not tag.find_parent(HTML_CONTENT_PARENTS):
            tag.decompose()

    def own_text(element, skip_tags) -> str:
        """Текст элемента без вложенных блоков из skip_tags и без комментариев разметки."""
        parts = []
        for child in element.children:
            if isinstance(child, bs4.Comment):
                continue
            if isinstance(child, bs4.NavigableString):
                parts.append(str(child))
            elif child.name not in skip_tags:
                parts.append(own_text(child, skip_tags))
        return " ".join(part.strip() for part in parts if part.strip())

    paragraphs = []
    for element in soup.find_all(HTML_BLOCK_TAGS):
        # Содержимое таблиц разбирается вместе с таблицей
        if element.name != "table" and element.find_parent("table"):
            continue
        # Абзац внутри пункта списка даёт текст через сам пункт, чтобы не потерять его номер
        if element.name not in ("li", "dd", "ul", "ol", "table") and element.find_parent(["li", "dd"]):
            continue

        if element.name == "table":
            # Строки вложенной таблицы относятся к ней, а не к внешней
            rows = [
                [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])]
                for row in element.find_all("tr") if row.find_parent("table") is element
            ]
            caption = element.find("caption")
            name = caption.get_text(" ", strip=True) if caption else f"{path.stem}, таблица {len(document.tables) + 1}"
            if len([row for row in rows if any(row)]) >= 2:
                _add_table(document, name, rows)
            else:
                # Таблица из одной строки (вёрстка, а не данные): её текст идёт обычным абзацем
                paragraphs.append(" ".join(cell for row in rows for cell in row if cell))
            continue

        if element.name in ("li", "dd"):
            # Пункт списка целиком, кроме вложенных списков и таблиц: они разбираются сами
            text = own_text(element, ["ul", "ol", "table"])
        else:
            # У блока с вложенными блоками берётся только его собственный текст:
            # вложенные блоки дадут свой текст сами
            text = own_text(element, HTML_BLOCK_TAGS + ["ul", "ol"])
        if not text:
            continue

        if element.name.startswith("h") and len(element.name) == 2:
            paragraphs.append(f"# {text}")
        elif element.name == "li":
            parent = element.find_parent(["ol", "ul"])
            if parent is not None and parent.name == "ol":
                number = len(element.find_previous_siblings("li")) + 1
                paragraphs.append(f"{number}. {text}")
            else:
                paragraphs.append(f"- {text}")
        else:
            paragraphs.append(text)

    document.text = _join_paragraphs(paragraphs)


def _join_paragraphs(paragraphs: List[str]) -> str:
    """
    Сборка текста из абзацев. Пункты одного списка идут строками подряд,
    остальные абзацы разделяются пустой строкой.
    """
    parts = []
    for paragraph in paragraphs:
        is_item = bool(LIST_ITEM.match(paragraph))
        previous_is_item = bool(parts) and bool(LIST_ITEM.match(parts[-1].split("\n")[-1]))
        if is_item and previous_is_item:
            parts[-1] += "\n" + paragraph
        else:
            parts.append(paragraph)
    return "\n\n".join(parts)


def _load_pdf(path: Path, document: Document, clean: bool):
    """Загрузка PDF: текст страниц без колонтитулов, номеров страниц и переносов слов."""
    pypdf = _require("pypdf", "pypdf", ".pdf")
    reader = pypdf.PdfReader(str(path))
    pages = [(page.extract_text() or "") for page in reader.pages]

    if not clean:
        document.text = "\n".join(pages)
        return

    page_lines = [[line.strip() for line in page.split("\n") if line.strip()] for page in pages]

    # Колонтитул - строка у верхнего или нижнего края страницы, которая повторяется
    # на большинстве страниц. Строки из середины страницы колонтитулом не считаются
    def edge_lines(page: List[str]) -> List[str]:
        return page[:PDF_EDGE_LINES] + page[-PDF_EDGE_LINES:]

    repeated = set()
    if len(page_lines) >= 3:
        counts = Counter(line for page in page_lines for line in set(edge_lines(page)))
        repeated = {line for line, count in counts.items() if count >= len(page_lines) * 0.6}

    lines = []
    for page_number, page in enumerate(page_lines, 1):
        # Сначала убираются колонтитулы и номера с подписью: номер без подписи ищется уже после них
        page = [
            line for position, line in enumerate(page)
            if not PAGE_LABEL.match(line)
            and not (line in repeated and (position < PDF_EDGE_LINES or position >= len(page) - PDF_EDGE_LINES))
        ]
        for position, line in enumerate(page):
            # Число без подписи - номер страницы, только если оно равно номеру этой страницы
            # и стоит первой или последней строкой. Иначе это значение из текста или таблицы
            bare = PAGE_BARE_NUMBER.match(line)
            if bare and int(bare.group(1)) == page_number and position in (0, len(page) - 1):
                continue
            lines.append(line)

    document.text = _join_pdf_lines(lines)


def _join_pdf_lines(lines: List[str]) -> str:
    """
    Сборка абзацев из строк PDF.

    В PDF строка обрывается по ширине страницы, а не по смыслу. Строки склеиваются
    в абзац, пока предыдущая не закончится концом предложения, а следующая
    не начнётся с заглавной буквы, цифры или пункта списка.
    """
    # Обычная ширина строки в этом файле: заголовок заметно короче неё
    long_lines = [len(line) for line in lines if len(line) > 20]
    usual_width = statistics.median(long_lines) if long_lines else 80

    paragraphs = []
    current = ""
    # Абзац пока состоит из одной строки файла: только такая строка может быть заголовком
    current_is_single_line = True

    for line in lines:
        # Маркер списка приводится к общему виду
        line = re.sub(r"^•\s+", "- ", line)

        if not current:
            current = line
            current_is_single_line = True
            continue

        # Перенос слова по слогам: "конвер-" + "сия". Частицы через дефис не склеиваются
        if current.endswith("-") and line[:1].islower():
            first_word = re.split(r"\W", line, maxsplit=1)[0]
            current = current + line if first_word in HYPHEN_PARTICLES else current[:-1] + line
            current_is_single_line = False
            continue

        ends_sentence = current[-1] in ".!?:;"
        starts_new = line[:1].isupper() or line[:1].isdigit()
        # Пункт списка всегда начинает новый абзац
        is_list_item = bool(PDF_LIST_ITEM.match(line))
        # Заголовок (в том числе с номером: "2. Установка") или последний пункт списка:
        # одна короткая строка без знака препинания в конце перед строкой с заглавной буквы
        looks_like_heading = (
            current_is_single_line
            and len(current) < usual_width * 0.6
            and current[-1] not in ".,;:!?"
            and line[:1].isupper()
        )

        if is_list_item or (ends_sentence and starts_new) or looks_like_heading:
            paragraphs.append(current)
            current = line
            current_is_single_line = True
        else:
            current += " " + line
            current_is_single_line = False

    if current:
        paragraphs.append(current)

    return _join_paragraphs(paragraphs)


def _load_docx(path: Path, document: Document):
    """Загрузка документа Word: абзацы, заголовки по стилю и таблицы, в порядке следования."""
    docx = _require("docx", "python-docx", ".docx")
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph

    source = docx.Document(str(path))
    paragraphs = []
    item_number = 0

    for element in source.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]

        if tag == "p":
            paragraph = Paragraph(element, source)
            text = paragraph.text.strip()
            if not text:
                continue

            style = (paragraph.style.name if paragraph.style is not None else "").lower()
            if style.startswith(("heading", "заголовок", "title", "название")):
                paragraphs.append(f"# {text}")
            elif "list" in style or "список" in style:
                # Номер пункта Word хранит отдельно от текста и в самом абзаце его нет:
                # для нумерованного списка номер восстанавливается счётчиком
                is_numbered = "number" in style or "нумер" in style
                item_number = item_number + 1 if is_numbered else 0
                if LIST_ITEM.match(text):
                    # Номер или маркер уже набран в тексте: не дублируем
                    paragraphs.append(text)
                elif is_numbered:
                    paragraphs.append(f"{item_number}. {text}")
                else:
                    paragraphs.append(f"- {text}")
                continue
            else:
                paragraphs.append(text)

            # Любой абзац, кроме пункта нумерованного списка, начинает нумерацию заново
            item_number = 0

        elif tag == "tbl":
            table = DocxTable(element, source)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            _add_table(document, f"{path.stem}, таблица {len(document.tables) + 1}", rows)

    document.text = _join_paragraphs(paragraphs)


def _load_csv(path: Path, document: Document):
    """Загрузка таблицы CSV; разделитель определяется по первой строке."""
    text = _read_text(path)
    first_line = text.split("\n", 1)[0]
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    if first_line.count("\t") > max(first_line.count(";"), first_line.count(",")):
        delimiter = "\t"

    # StringIO, а не список строк: ячейка в кавычках может занимать несколько строк
    rows = [[_cell_text(cell) for cell in row] for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)]
    _add_table(document, path.stem, rows)


def _load_xlsx(path: Path, document: Document):
    """Загрузка книги Excel: каждый лист - отдельная таблица."""
    openpyxl = _require("openpyxl", "openpyxl", ".xlsx")
    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)

    try:
        for sheet in workbook.worksheets:
            rows = [[_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
            _add_table(document, f"{path.stem}, лист {sheet.title}", rows)
    finally:
        workbook.close()
